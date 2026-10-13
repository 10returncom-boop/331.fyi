"""Regenerate Excel deliverables with current 23-page structure."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = r"D:\WWW\325_pub\自建主題目錄"
DOCS = ROOT + r"\docs"

# === 1. 網站地圖.xlsx ===
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "網站地圖"
header_fill = PatternFill("solid", fgColor="1F7A4D")
header_font = Font(bold=True, color="FFFFFF")
ws.append(["頁面路徑", "類型", "所屬主題", "標題", "lastmod"])
for c in ws[1]:
    c.fill = header_fill; c.font = header_font

pages = [
    ("index.html", "首頁", "-", "營養與生活：大白話營養知識庫", "2026-10-13"),
    ("about.html", "關於", "-", "關於我們", "2026-10-12"),
    ("404.html", "404", "-", "找不到頁面", "2026-10-12"),
    ("nutrition-basics/index.html", "Hub", "營養基礎", "三大營養素基礎", "2026-10-12"),
    ("nutrition-basics/style-protein.html", "文章", "營養基礎", "蛋白質完整指南", "2026-10-12"),
    ("nutrition-basics/style-vitamin.html", "文章", "營養基礎", "維生素完整指南", "2026-10-12"),
    ("diet-disease/index.html", "Hub", "飲食與疾病", "飲食與慢性病", "2026-10-12"),
    ("diet-disease/style-diabetes.html", "文章", "飲食與疾病", "糖尿病飲食", "2026-10-12"),
    ("diet-disease/style-hypertension.html", "文章", "飲食與疾病", "高血壓飲食", "2026-10-12"),
    ("life-stage/index.html", "Hub", "各族群營養", "各人生階段營養", "2026-10-12"),
    ("life-stage/style-senior.html", "文章", "各族群營養", "銀髮族營養", "2026-10-12"),
    ("life-stage/style-pregnant.html", "文章", "各族群營養", "孕婦營養", "2026-10-12"),
    ("food-guide/index.html", "Hub", "食材與標示", "食材選購指南", "2026-10-12"),
    ("food-guide/style-label.html", "文章", "食材與標示", "營養標示解讀", "2026-10-12"),
    ("food-guide/style-superfood.html", "文章", "食材與標示", "超級食物", "2026-10-12"),
    ("lifestyle/index.html", "Hub", "生活型態", "減重與斷食", "2026-10-12"),
    ("lifestyle/style-weight.html", "文章", "生活型態", "減重方法", "2026-10-12"),
    ("lifestyle/style-fasting.html", "文章", "生活型態", "間歇性斷食", "2026-10-12"),
    ("courses/index.html", "Hub", "線上課程", "線上營養課程總覽", "2026-10-13"),
    ("courses/style-course-diabetes.html", "課程", "線上課程", "糖尿病低GI實戰班", "2026-10-13"),
    ("courses/style-course-weight.html", "課程", "線上課程", "減重不挨餓", "2026-10-13"),
    ("courses/style-course-basics.html", "課程", "線上課程", "21天飲食健檢", "2026-10-13"),
    ("courses/style-course-senior.html", "課程", "線上課程", "銀髮族防肌少飲食", "2026-10-13"),
]
for p in pages:
    ws.append(p)
for col, w in zip("ABCDE", [40, 8, 14, 30, 14]):
    ws.column_dimensions[col].width = w
wb.save(DOCS + r"\網站地圖.xlsx")
print("網站地圖.xlsx: 23 pages")

# === 2. 網站變數與導覽表.xlsx ===
wb2 = openpyxl.Workbook()
ws1 = wb2.active
ws1.title = "網站變數"
ws1.append(["變數名", "值", "說明"])
for c in ws1[1]:
    c.fill = header_fill; c.font = header_font
vars_data = [
    ("site_name", "營養與生活", "網站名稱"),
    ("domain", "https://nutrition.example.com", "正式網域（部署時替換）"),
    ("color_primary", "#1f7a4d", "主題色深綠"),
    ("mode_default", "light", "預設白天模式"),
    ("og_level", "9", "Open Graph 圖等級"),
    ("generator", "build.py + data/*.json", "資料驅動生成器"),
    ("encoding", "UTF-8", "編碼"),
    ("stack", "HTML5 + CSS3 + Vanilla JS", "無框架"),
]
for v in vars_data:
    ws1.append(v)
ws1.column_dimensions["A"].width = 18
ws1.column_dimensions["B"].width = 35
ws1.column_dimensions["C"].width = 30

ws2 = wb2.create_sheet("網站導覽")
ws2.append(["主選單", "子選單1", "子選單2", "連結"])
for c in ws2[1]:
    c.fill = header_fill; c.font = header_font
nav = [
    ("首頁", "", "", "index.html"),
    ("營養基礎", "蛋白質", "", "nutrition-basics/style-protein.html"),
    ("營養基礎", "維生素", "", "nutrition-basics/style-vitamin.html"),
    ("飲食與疾病", "糖尿病", "", "diet-disease/style-diabetes.html"),
    ("飲食與疾病", "高血壓", "", "diet-disease/style-hypertension.html"),
    ("各族群營養", "銀髮族", "", "life-stage/style-senior.html"),
    ("各族群營養", "孕婦", "", "life-stage/style-pregnant.html"),
    ("食材與標示", "營養標示", "", "food-guide/style-label.html"),
    ("食材與標示", "超級食物", "", "food-guide/style-superfood.html"),
    ("生活型態", "減重", "", "lifestyle/style-weight.html"),
    ("生活型態", "斷食", "", "lifestyle/style-fasting.html"),
    ("線上課程", "糖尿病實戰", "", "courses/style-course-diabetes.html"),
    ("線上課程", "減重班", "", "courses/style-course-weight.html"),
    ("線上課程", "營養基礎", "", "courses/style-course-basics.html"),
]
for n in nav:
    ws2.append(n)
for col, w in zip("ABCD", [14, 14, 10, 40]):
    ws2.column_dimensions[col].width = w

ws3 = wb2.create_sheet("標籤清單")
ws3.append(["標籤", "頁數"])
for c in ws3[1]:
    c.fill = header_fill; c.font = header_font
tags = [("蛋白質", 2), ("維生素", 1), ("糖尿病", 2), ("高血壓", 1), ("銀髮族", 2),
        ("孕婦", 1), ("營養標示", 1), ("減重", 2), ("斷食", 1), ("線上課程", 4)]
for t in tags:
    ws3.append(t)
wb2.save(DOCS + r"\網站變數與導覽表.xlsx")
print("網站變數與導覽表.xlsx: 3 sheets")
